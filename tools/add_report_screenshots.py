"""Añade capturas reales al informe existente, conservando las ediciones del usuario."""
from pathlib import Path
import argparse
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--source', type=Path, default=ROOT / 'output/word/seccion2-rest-api-todo-ly.docx')
args = parser.parse_args()
doc = Document(args.source)
if any(p.text == 'Capturas del trabajo realizado' for p in doc.paragraphs):
    raise SystemExit('El documento ya contiene el anexo de capturas; use una copia anterior como --source.')

section = doc.add_section(WD_SECTION_START.NEW_PAGE)
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = Pt(841.89), Pt(595.276)
section.left_margin = section.right_margin = Pt(36)
section.top_margin, section.bottom_margin = Pt(30), Pt(36)
section.footer_distance = Pt(17)

items = [
    ('01-resultados-cucumber.jpg', 'Capturas del trabajo realizado',
     'Figura 1  Resultado del caso TC API 001',
     'El reporte generado por Cucumber registra 1 prueba, 0 fallos, 0 errores y los 11 pasos con estado passed. '
     'La ejecución corresponde al 20 de octubre de 2026; la captura se obtuvo el 24 de octubre de 2026.',
     'docs/evidence/cucumber.xml', (15, 170, 30, 0)),
    ('02-escenario-gherkin.jpg', 'Escenario automatizado en Gherkin',
     'Figura 2  Definición del flujo y sus verificaciones',
     'El archivo feature contiene las etiquetas de ejecución y los pasos de creación de usuario, obtención del token, '
     'creación, consulta, actualización y eliminación del proyecto.',
     'src/test/resources/features/project_token.feature', (15, 170, 30, 75)),
    ('03-cliente-token.jpg', 'Implementación de la autenticación',
     'Figura 3  Uso del encabezado Token con Rest Assured',
     'El método token valida que el valor no esté vacío y lo envía mediante header("Token", token). '
     'El método send registra la operación, la autenticación, el código HTTP y el tiempo de respuesta.',
     'src/test/java/client/ApiClient.java', (15, 0, 330, 35)),
    ('04-dependencias-gradle.jpg', 'Tecnologías y dependencias del proyecto',
     'Figura 4  Configuración de Java Cucumber y Rest Assured',
     'build.gradle aplica el plugin java y declara Cucumber 8.0.4, Rest Assured 6.0.1 y su validador JSON Schema. '
     'Las versiones coinciden con las utilizadas en el caso de prueba documentado.',
     'build.gradle', (15, 170, 30, 92)),
]

for i, (filename, heading, caption, description, source, crop) in enumerate(items):
    if i:
        doc.add_page_break()
    p = doc.add_paragraph(heading, style='Heading 1')
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = True
    p = doc.add_paragraph(description)
    p.paragraph_format.line_spacing = Pt(13)
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.size = Pt(9.5)
    path = ROOT / 'docs/evidence/screenshots' / filename
    with Image.open(path) as image:
        w, h = image.size
    left, top, right, bottom = crop
    visible_w, visible_h = w-left-right, h-top-bottom
    scale = min(769.89/visible_w, 410/visible_h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1
    shape = p.add_run().add_picture(str(path), width=Pt(visible_w*scale), height=Pt(visible_h*scale))
    # Recorte no destructivo en Word: las capturas originales permanecen completas.
    blipfill = shape._inline.graphic.graphicData.pic.blipFill
    src = OxmlElement('a:srcRect')
    for name, value, total in [('l', left, w), ('t', top, h), ('r', right, w), ('b', bottom, h)]:
        src.set(name, str(round(value/total*100000)))
    blipfill.insert(1, src)
    shape._inline.docPr.set('descr', caption + '. ' + description)
    p = doc.add_paragraph(caption)
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        run.bold = True
        run.font.size = Pt(9)
    p = doc.add_paragraph('Fuente: ')
    p.paragraph_format.space_after = Pt(0)
    relationship = p.part.relate_to(
        'https://github.com/nohpedro/API-Testing/blob/77bc799349d523fdc08f889c5d20131a86a93502/' + source,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), relationship)
    run = p.add_run(source + ' en GitHub')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string('007C83')
    link.append(run._r)
    p._p.append(link)

output = ROOT / 'output/word/seccion2-rest-api-todo-ly.docx'
doc.save(output)
assert len(Document(output).inline_shapes) >= 4
print(output)
