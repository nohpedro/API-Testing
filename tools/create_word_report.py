"""Crea una versión Word editable reutilizando el contenido del informe PDF."""
import contextlib
import io
import runpy
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import html
from reportlab.platypus import Paragraph, Table, Preformatted, PageBreak, Spacer

ROOT = Path(__file__).resolve().parents[1]
with patch("reportlab.platypus.SimpleDocTemplate.build"), contextlib.redirect_stdout(io.StringIO()):
    source = runpy.run_path(str(ROOT / "tools/create_report.py"))

doc = Document()
section = doc.sections[0]
section.page_width = Pt(595.276)
section.page_height = Pt(841.89)
section.top_margin = Pt(38)
section.bottom_margin = Pt(48)
section.left_margin = section.right_margin = Pt(44)
section.footer_distance = Pt(20)

for name in ["Normal", "Title", "Heading 1", "Heading 2", "Subtitle"]:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(8)
doc.styles["Normal"].font.size = Pt(10)
doc.styles["Normal"].paragraph_format.line_spacing = Pt(14)
doc.styles["Title"].font.size = Pt(27)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.line_spacing = Pt(31)
doc.styles["Heading 1"].font.size = Pt(18)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(0)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(12)
doc.styles["Heading 2"].font.size = Pt(11)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(6)


def xml_element(name, **attributes):
    element = OxmlElement(name)
    for key, value in attributes.items():
        element.set(qn("w:" + key), str(value))
    return element


def rich_text(paragraph, markup, force_white=False):
    root = html.fragment_fromstring(markup.replace("<link ", "<a ").replace("</link>", "</a>"),
                                   create_parent="div")

    def emit(text, bold=False, url=None):
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold = bold
        if force_white:
            run.font.color.rgb = RGBColor(255, 255, 255)
        if url:
            rel = paragraph.part.relate_to(
                url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), rel)
            run.font.color.rgb = RGBColor.from_string("007C83")
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)

    def walk(node, bold=False, url=None):
        bold = bold or node.tag in ("b", "strong")
        url = node.get("href") if node.tag in ("link", "a") else url
        emit(node.text, bold, url)
        for child in node:
            if child.tag == "br":
                paragraph.add_run().add_break()
            else:
                walk(child, bold, url)
            emit(child.tail, bold, url)
    walk(root)


for item in source["story"]:
    if isinstance(item, PageBreak):
        doc.add_page_break()
    elif isinstance(item, Spacer):
        # Las tablas ya tienen márgenes y el siguiente párrafo aporta separación.
        pass
    elif isinstance(item, Paragraph):
        kind = item.style.name
        if kind == "Kicker":
            continue
        markup = item.text
        style = {"TitleReport": "Title", "Section": "Heading 1", "Sub": "Heading 2"}.get(kind, "Normal")
        if style.startswith("Heading"):
            markup = markup.replace("Caso de prueba: pasos y resultados", "Pasos y resultados del caso de prueba")
        if kind == "TitleReport":
            markup = "Pruebas de proyectos en Todo ly con token"
        paragraph = doc.add_paragraph(style=style)
        rich_text(paragraph, markup)
        if style in ("Title", "Heading 1", "Heading 2"):
            for run in paragraph.runs:
                run.bold = True
        if kind == "SmallReport":
            paragraph.paragraph_format.line_spacing = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                run.font.size = Pt(8.7)
    elif isinstance(item, Preformatted):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = Pt(11.5)
        paragraph.paragraph_format.space_after = Pt(10)
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run("\n".join(item.lines))
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
    elif isinstance(item, Table):
        table = doc.add_table(rows=len(item._cellvalues), cols=len(item._colWidths))
        table.autofit = False
        for column, width in zip(table.columns, item._colWidths):
            column.width = Pt(width)
        properties = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(xml_element("w:" + edge, val="single", sz="4", color="D9D9D9"))
        properties.append(borders)
        margins = OxmlElement("w:tblCellMar")
        for edge, value in [("top", 120), ("bottom", 120), ("left", 150), ("right", 150)]:
            margins.append(xml_element("w:" + edge, w=value, type="dxa"))
        properties.append(margins)
        for row_index, values in enumerate(item._cellvalues):
            row = table.rows[row_index]
            row._tr.get_or_add_trPr().append(xml_element("w:cantSplit"))
            if row_index == 0:
                row._tr.get_or_add_trPr().append(xml_element("w:tblHeader"))
            for column_index, value in enumerate(values):
                cell = row.cells[column_index]
                cell.width = Pt(item._colWidths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                fill = "132B45" if row_index == 0 else ("EEF5F7" if row_index % 2 == 0 else "FFFFFF")
                cell._tc.get_or_add_tcPr().append(xml_element("w:shd", fill=fill))
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(12)
                if item._colWidths[column_index] < 85:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rich_text(paragraph, value.text, force_white=row_index == 0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.7)
                    if row_index == 0:
                        run.bold = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("TC-API-001  |  Página ")
footer._p.append(xml_element("w:fldSimple", instr="PAGE"))
for run in footer.runs:
    run.font.name = "Arial"
    run.font.size = Pt(8)

doc.core_properties.title = "Pruebas de proyectos en Todo ly con token"
doc.core_properties.subject = "Sección 2 REST API con Java Gradle Cucumber y Rest Assured"
doc.core_properties.author = "nohpedro"
for element in [doc.styles.element, doc.element]:
    for border in element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
output = ROOT / "output/word/seccion2-rest-api-todo-ly.docx"
output.parent.mkdir(parents=True, exist_ok=True)
doc.save(output)
print(output)
